"""The spreadsheet importers, on small workbooks made here: py -m unittest"""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from cli import importers, model, storage

YELLOW = PatternFill(fill_type="solid", fgColor="00FFF2CC")


def site(rid, url, ages=(13, 18)):
    return {"id": rid, "kind": "website", "title": "A site", "url": url, "languages": ["en"],
            "classes": ["lower-secondary-beginner", "upper-secondary-middle"],
            "ages": {"min": ages[0], "max": ages[1]}, "strands": [], "lenses": [],
            "role": "religious", "rights": "pending", "status": "to_verify",
            "retrieved": "2026-10-13"}


class TopicSheetTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "topics.xlsx"
        wb = Workbook()
        wb.remove(wb.active)
        rows = {
            "S4 (Age 17-18)": [
                (1, "Introduction", "Said Nursi, The Words – The First Word (Bismillah) (risaleinur.com)",
                 "https://risaleinur.com/risale-i-nur-collection/75-words/3655-01.html", False),
            ],
            "S6 (Age 18-19)": [
                (5, "The Heart and Mind Connection",
                 "Gülen, Emerald Hills of the Heart (Key Concepts in Sufism) – 'Aql (Reason) (fgulen.com)",
                 "https://fgulen.com/en/fethullah-gulens-works/key-concepts-in-the-practice-of-sufism-3/reason", False),
                (6, "Patience", "Related reading: Gülen, Emerald Hills of the Heart (Key Concepts in Sufism) – Sabr (fgulen.com)",
                 "https://fgulen.com/en/fethullah-gulens-works/key-concepts-in-the-practice-of-sufism-1/sabr", True),
                (7, "Same chapter again", "Said Nursi, The Words – The First Word (Bismillah) (risaleinur.com)",
                 "https://risaleinur.com/risale-i-nur-collection/75-words/3655-01.html", False),
                (8, "Elsewhere", "Green Dome Foundation – A Story (greendome.org)",
                 "https://greendome.org/a-story/", False),
            ],
        }
        for title, entries in rows.items():
            ws = wb.create_sheet(title)
            ws.append(["Week", "Topic", "Source"])
            for week, topic, source, url, related in entries:
                ws.append([week, topic, source])
                cell = ws.cell(row=ws.max_row, column=3)
                cell.hyperlink = url
                if related:
                    cell.fill = YELLOW
            ws.append([None, "Click a source to open it.", None])
        wb.save(self.path)
        self.cat = storage.empty()
        self.cat["resources"] = [site("r-fg", "https://www.fgulen.com/en/"),
                                 site("r-rn", "https://risaleinur.com/risale-i-nur-collection.html")]

    def tearDown(self):
        self.tmp.cleanup()

    def run_import(self):
        return importers.import_topics(self.cat, importers.topic_rows(str(self.path)), "2026-10-13")

    def test_series_map_to_the_sites_ages(self):
        ages = sorted({r["age"] for r in importers.topic_rows(str(self.path))})
        self.assertEqual(ages, [16, 18])  # S4 is labelled 17-18 but is age 16

    def test_a_hint_names_the_book_not_the_author_publisher_or_site(self):
        self.run_import()
        hints = {(h["age"], h["topic"]): h for h in self.cat["topicHints"]}
        self.assertEqual(hints[(18, 5)]["hint"], "Emerald Hills of the Heart (Key Concepts in Sufism) – 'Aql (Reason)")
        self.assertEqual(hints[(18, 5)]["fromSheet"], "S6 week 5")
        self.assertEqual(hints[(18, 8)]["hint"], "A Story")
        self.assertEqual(hints[(18, 6)]["hint"],
                         "Emerald Hills of the Heart (Key Concepts in Sufism) – Sabr (related reading)")

    def test_only_a_catalogued_volume_gets_its_catalogue_id(self):
        self.run_import()
        ids = {h["topic"]: h["catalogueId"] for h in self.cat["topicHints"] if h["age"] == 18}
        self.assertIsNone(ids[5])            # volume 3 is not in the site's catalogue
        self.assertEqual(ids[6], "s-6b15280cee")

    def test_pages_become_items_of_their_site_and_widen_across_ages(self):
        changed, _, skipped = self.run_import()
        by_url = {i["url"]: i for i in self.cat["items"]}
        words = by_url["https://risaleinur.com/risale-i-nur-collection/75-words/3655-01.html"]
        self.assertEqual((words["resourceId"], words["kind"], words["ages"]),
                         ("r-rn", "chapter", {"min": 16, "max": 18}))
        self.assertEqual(by_url["https://fgulen.com/en/fethullah-gulens-works/key-concepts-in-the-practice-of-sufism-3/reason"]["resourceId"], "r-fg")
        self.assertEqual(skipped, ["S6 week 8: no resource for greendome.org yet"])
        self.assertEqual(model.check(self.cat), [])

    def test_running_it_again_changes_nothing(self):
        self.run_import()
        changed, _, _ = self.run_import()
        self.assertEqual(changed, [])
        self.assertEqual(len(self.cat["topicHints"]), 5)


class VideoSheetTest(unittest.TestCase):
    CHANNEL = "https://www.youtube.com/@%C3%87%C4%B1narMedya"

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "videos.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append([None, None, "VAAZ İSMİ", "SÜRE", "KAYNAK", "LİNK"])
        for number, title, minutes, vid in [(1, "TEVHID DELILLERI-1", 16, "-hpH6I13U18"),
                                            (2, "GONE AWAY", 10, "AAAAAAAAAAA"),
                                            (3, "OTHER CHANNEL", 5, "BBBBBBBBBBB"),
                                            (4, "TEVHID DELILLERI-1", 16, "-hpH6I13U18"),
                                            (5, "Example Academy talk", 9, "CCCCCCCCCCC")]:
            ws.append([None, number, title, minutes, "CINAR MEDYA ENGLISH SUBTITLES", "link"])
            ws.cell(row=ws.max_row, column=6).hyperlink = \
                f"https://www.youtube.com/watch?v={vid}&list=PLx&ab_channel=x"
        ws.append([None, None, None, "=SUM(D3:D502)", None, None])
        wb.save(self.path)
        self.cat = storage.empty()
        channel = site("r-ch", "https://www.youtube.com/@çınarmedya/")
        channel.update(kind="channel", languages=["tr"], role="religious-supporting")
        self.cat["resources"] = [channel]
        here = {"status": 200, "title": "t", "channel": "Çınar Medya", "channelUrl": self.CHANNEL}
        self.found = {"-hpH6I13U18": here, "AAAAAAAAAAA": {"status": 404},
                      "BBBBBBBBBBB": dict(here, channel="Other", channelUrl="https://www.youtube.com/@other"),
                      "CCCCCCCCCCC": here}

    def tearDown(self):
        self.tmp.cleanup()

    def test_videos_land_under_their_channel_once_and_the_rest_is_reported(self):
        from tests.test_catalogue import BLOCKED_NAME, fingerprint
        from cli.names import BlockedNames
        blocked = BlockedNames({"version": 1, "names": [fingerprint(BLOCKED_NAME)]})
        rows = list(importers.video_rows(str(self.path)))
        changed, skipped = importers.import_videos(self.cat, rows, self.found, "2026-10-13", blocked)
        self.assertEqual(len(changed), 1)
        item = self.cat["items"][0]
        self.assertEqual((item["resourceId"], item["url"], item["minutes"], item["subtitles"],
                          item["youtubeId"], item["languages"]),
                         ("r-ch", "https://www.youtube.com/watch?v=-hpH6I13U18", 16, ["en"],
                          "-hpH6I13U18", ["tr"]))
        self.assertNotIn("content", item)  # only someone who watched it may say what is in it
        self.assertIn("row 2: removed or private (AAAAAAAAAAA)", skipped)
        self.assertIn("row 5: its title carries a blocked name", skipped)
        self.assertIn("1 video(s): channel 'Other' is not catalogued yet: https://www.youtube.com/@other",
                      skipped)
        self.assertEqual(model.check(self.cat), [])
        again, _ = importers.import_videos(self.cat, rows, self.found, "2026-10-13", blocked)
        self.assertEqual(again, [])


class DiscoverTest(unittest.TestCase):
    def test_titles_from_addresses(self):
        from cli.text import title_from_url
        self.assertEqual(title_from_url("https://x.org/nebevi-bir-yol-genis-tabanli-hareket-2/"),
                         "Nebevi bir yol genis tabanli hareket 2")
        self.assertEqual(title_from_url("https://x.org/75-words/3663-03-the-third-word.html"),
                         "03 the third word")

    def test_only_what_holds_the_words_comes_in_once(self):
        from tests.test_catalogue import BLOCKED_NAME, fingerprint
        from cli.names import BlockedNames
        blocked = BlockedNames({"version": 1, "names": [fingerprint(BLOCKED_NAME)]})
        cat = storage.empty()
        channel = site("r-ch", "https://www.youtube.com/@x")
        channel.update(kind="channel", languages=["tr"])
        cat["resources"] = [channel]
        entries = [
            {"url": "https://www.youtube.com/watch?v=AAAAAAAAAAA", "title": "Arkadaş seçimi",
             "kind": "video", "youtubeId": "AAAAAAAAAAA", "minutes": 12},
            {"url": "https://www.youtube.com/watch?v=BBBBBBBBBBB", "title": "Ramazan",
             "kind": "video", "youtubeId": "BBBBBBBBBBB", "minutes": 8},
            {"url": "https://www.youtube.com/watch?v=CCCCCCCCCCC", "title": "Example Academy arkadas gunu",
             "kind": "video", "youtubeId": "CCCCCCCCCCC", "minutes": 3},
        ]
        words = ["arkadas"]  # as normalise() writes it: accents and case aside
        matching = [e for e in entries if "arkadas" in __import__("cli.names").names.normalise(e["title"])]
        added, skipped = importers.discover(cat, channel, matching, words, "2026-10-15", blocked)
        self.assertEqual([i["title"] for i in cat["items"]], ["Arkadaş seçimi"])
        self.assertEqual(skipped, 1)
        self.assertEqual((cat["items"][0]["minutes"], cat["items"][0]["languages"]), (12, ["tr"]))
        self.assertEqual(model.check(cat), [])
        self.assertEqual(importers.discover(cat, channel, matching, words, "2026-10-15", blocked)[0], [])


if __name__ == "__main__":
    unittest.main()
