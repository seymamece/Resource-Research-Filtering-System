"""What comes out for people: the review sheet there and back, and a topic's pack."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from cli import model, packs, review, storage
from tests.test_sift import resource, video


def catalogue():
    cat = storage.empty()
    cat["resources"] = [resource("r-a", role="religious")]
    cat["items"] = [video("v1", "r-a", 12), video("v2", "r-a", 9), video("v3", "r-a", 7)]
    for pid, where in (("p-1", "01:00-01:20"), ("p-2", "02:00-02:20"), ("p-3", "03:00-03:20"),
                       ("p-4", "04:00-04:20")):
        cat["passages"].append({"id": pid, "itemId": "v1", "where": where, "topics": [{"age": 18, "topic": 4}],
                                "use": "faith", "status": "proposed"})
    return cat


class ReviewSheetTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "review.xlsx"
        self.cat = catalogue()

    def tearDown(self):
        self.tmp.cleanup()

    def fill(self, passages: dict, videos: dict):
        wb = load_workbook(self.path)
        for sheet, values, columns in (("Passages", passages, review.PASSAGE_COLUMNS),
                                       ("Videos", videos, review.VIDEO_COLUMNS)):
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=2):
                for name, value in values.get(row[0].value, {}).items():
                    row[columns.index(name)].value = value
        wb.save(self.path)

    def test_only_rows_that_name_a_person_are_applied(self):
        self.assertEqual(review.write(self.path, self.cat, {(18, 4): {"title": "Daily Worship"}}, 18), (4, 3))
        self.fill({"p-1": {"Decision": "accept", "Reviewed by": "Ayse"},
                   "p-2": {"Decision": "reject", "Reviewed by": "Ayse"},
                   "p-3": {"Decision": "accept"},
                   "p-4": {"Decision": "maybe", "Reviewed by": "Ayse"}},
                  {"v1": {"Content": "talk, music", "Watched by": "Ali"},
                   "v2": {"Content": "faces"},
                   "v3": {"Content": "none of these", "Watched by": "Ali"}})
        changed, report = review.read(self.path, self.cat, "2026-10-15")
        passages = {p["id"]: p for p in self.cat["passages"]}
        self.assertEqual(passages["p-1"]["acceptedBy"], {"by": "Ayse", "on": "2026-10-15"})
        self.assertNotIn("p-2", passages)
        self.assertEqual(passages["p-3"]["status"], "proposed")
        self.assertEqual(passages["p-4"]["status"], "proposed")
        items = {i["id"]: i for i in self.cat["items"]}
        self.assertEqual((items["v1"]["content"], items["v1"]["watchedBy"]["by"]), (["talk", "music"], "Ali"))
        self.assertNotIn("watchedBy", items["v2"])
        self.assertEqual(items["v3"]["content"], [])
        self.assertEqual(sorted(changed), ["p-1", "v1", "v3"])
        self.assertTrue(any("p-3: accept without a name" in line for line in report))
        self.assertTrue(any("p-4: decision 'maybe'" in line for line in report))
        self.assertEqual(model.check(self.cat), [])

    def test_a_sheet_whose_columns_were_changed_is_refused(self):
        review.write(self.path, self.cat, {}, None)
        wb = load_workbook(self.path)
        wb["Passages"]["J1"].value = "Choice"
        wb.save(self.path)
        with self.assertRaises(ValueError):
            review.read(self.path, self.cat)


class PackTest(unittest.TestCase):
    def test_a_pack_lists_hint_passages_and_marks_what_cannot_reach_the_site(self):
        cat = catalogue()
        cat["passages"][0]["quote"] = "Worship gives peace and comfort to the heart."
        cat["items"][0]["youtubeId"] = "v1"
        cat["topicHints"] = [{"id": "h-1", "age": 18, "topic": 4, "hint": "The Third Word", "catalogueId": None}]
        corpus = packs.Corpus(cat)
        page = packs.build(cat, (18, 4), {"title": "Daily Worship", "science": ""}, "weak",
                           usable=set(), corpus=corpus)
        self.assertIn("# 18:4 Daily Worship", page)
        self.assertIn("- The Third Word", page)
        self.assertIn("> Worship gives peace and comfort to the heart.", page)
        self.assertIn("*not for the site*", page)
        self.assertIn("&t=60s", page)  # a video passage links to its moment


if __name__ == "__main__":
    unittest.main()
