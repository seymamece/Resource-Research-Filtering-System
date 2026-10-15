"""The owner's review sheet: proposed passages and unwatched videos out to Excel, a person's
decisions back in.

Rule 3 is why it has a "Reviewed by" column: nothing is accepted, and no video is marked
watched, unless the row names the person who did it. A row with a decision but no name is
reported and left as it was. Everything else in the sheet is for reading; only the columns
with a light yellow header are read back.
"""

from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import curriculum, vocab
from .packs import suits

FILL_IN = PatternFill(fill_type="solid", fgColor="FFF2CC")
HEADER = Font(bold=True)
DECISIONS = ("accept", "reject")

PASSAGE_COLUMNS = ["id", "Topic", "Topic title", "Use", "Quote", "Where", "Source", "Role",
                   "Link", "Decision", "Reviewed by", "Note"]
VIDEO_COLUMNS = ["id", "Title", "Minutes", "Subtitles", "Channel (private)", "Link",
                 "Content", "Watched by", "Note"]


def _sheet(wb, title, columns, rows, fill_in, widths):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = HEADER
        if cell.value in fill_in:
            cell.fill = FILL_IN
    for row in rows:
        ws.append(row)
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    return ws


def write(path, cat: dict, site_topics: dict, class_id: str | None = None,
          placement: dict | None = None):
    """Proposed passages (for the class's topics, if a class is given) in week order, and the
    unwatched videos that suit the class."""
    placement = placement or {}
    items = {i["id"]: i for i in cat["items"]}
    resources = {r["id"]: r for r in cat["resources"]}
    wb = Workbook()
    wb.remove(wb.active)

    def in_class(topic):
        return class_id is None or placement.get(topic, {}).get("class") == class_id

    def order(topic):
        week = placement.get(topic, {}).get("week", 99)
        return (placement.get(topic, {}).get("class", "~"), week, curriculum.sort_key(topic))

    rows = []
    for p in cat["passages"]:
        if p["status"] != "proposed":
            continue
        topics = [curriculum.key_of(t) for t in p["topics"]]
        mine = [t for t in topics if in_class(t)]
        if not mine:
            continue
        item = items.get(p["itemId"], {})
        resource = resources.get(item.get("resourceId"), {})
        rows.append([p["id"], ", ".join(topics), site_topics.get(mine[0], {}).get("title", ""),
                     p["use"], p.get("quote", ""), p["where"], item.get("title", ""),
                     resource.get("role", ""), item.get("url", ""), None, None, None])
    rows.sort(key=lambda r: order(r[1].split(", ")[0]))
    ws = _sheet(wb, "Passages", PASSAGE_COLUMNS, rows, {"Decision", "Reviewed by", "Note"},
                {"A": 12, "B": 8, "C": 30, "D": 10, "E": 60, "F": 22, "G": 34, "H": 14, "I": 30,
                 "J": 11, "K": 16, "L": 30})
    choice = DataValidation(type="list", formula1=f'"{",".join(DECISIONS)}"', allow_blank=True)
    ws.add_data_validation(choice)
    choice.add(f"J2:J{max(2, len(rows) + 1)}")

    videos = []
    for i in cat["items"]:
        if i.get("kind") != "video" or "watchedBy" in i:
            continue
        if class_id and not suits(i, resources.get(i["resourceId"], {}),
                                  {"class": class_id, "ages": curriculum.class_info(class_id)["ages"]}):
            continue
        videos.append([i["id"], i["title"], i.get("minutes"), ", ".join(i.get("subtitles", [])),
                       i.get("channelPrivate", ""), i["url"], None, None, None])
    _sheet(wb, "Videos", VIDEO_COLUMNS, videos, {"Content", "Watched by", "Note"},
           {"A": 12, "B": 50, "C": 9, "D": 10, "E": 26, "F": 44, "G": 30, "H": 16, "I": 30})

    guide = wb.create_sheet("How to fill in")
    for line in [
        "Passages: choose accept or reject in Decision, and write your name in Reviewed by.",
        "A row without a name is not read back. Reject removes the passage from the catalogue.",
        "Videos: after watching a video to the end, write your name in Watched by, and in",
        f"Content list what it has, separated by commas: {', '.join(vocab.VIDEO_CONTENT)}.",
        "Write 'none of these' when it has none of them. drawn-prophet means it is never usable.",
        "Only the yellow columns are read back; change nothing else.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 100
    wb.save(path)
    return len(rows), len(videos)


def read(path, cat: dict, today: str | None = None):
    """Applies the decisions in a filled-in sheet. Returns (changed ids, report lines)."""
    today = today or date.today().isoformat()
    wb = load_workbook(path)
    passages = {p["id"]: p for p in cat["passages"]}
    items = {i["id"]: i for i in cat["items"]}
    changed, report, rejected = [], [], set()

    def rows(name, columns):
        if name not in wb.sheetnames:
            return
        ws = wb[name]
        header = [c.value for c in ws[1]]
        if header != columns:
            raise ValueError(f"the '{name}' sheet's columns were changed; expected {columns}")
        for row in ws.iter_rows(min_row=2, values_only=True):
            yield dict(zip(columns, row))

    for row in rows("Passages", PASSAGE_COLUMNS):
        decision = str(row["Decision"] or "").strip().lower()
        if not decision:
            continue
        who = str(row["Reviewed by"] or "").strip()
        pid = row["id"]
        if pid not in passages:
            report.append(f"{pid}: no such passage any more")
        elif decision not in DECISIONS:
            report.append(f"{pid}: decision '{decision}' is not accept or reject")
        elif not who:
            report.append(f"{pid}: {decision} without a name in Reviewed by; left as it was")
        elif decision == "accept":
            passages[pid]["status"] = "accepted"
            passages[pid]["acceptedBy"] = {"by": who, "on": today}
            changed.append(pid)
        else:
            rejected.add(pid)
            report.append(f"{pid}: rejected by {who}, removed")
    if rejected:
        cat["passages"] = [p for p in cat["passages"] if p["id"] not in rejected]

    for row in rows("Videos", VIDEO_COLUMNS):
        who = str(row["Watched by"] or "").strip()
        content = str(row["Content"] or "").strip().lower()
        vid = row["id"]
        if not who and not content:
            continue
        if vid not in items:
            report.append(f"{vid}: no such video any more")
            continue
        if not who:
            report.append(f"{vid}: content given without a name in Watched by; left as it was")
            continue
        values = [] if content in ("", "none of these", "none") else [v.strip() for v in content.split(",")]
        unknown = [v for v in values if v not in vocab.VIDEO_CONTENT]
        if unknown:
            report.append(f"{vid}: {', '.join(unknown)} is not one of {', '.join(vocab.VIDEO_CONTENT)}")
            continue
        if not content:
            report.append(f"{vid}: watched by {who} but Content is empty; write 'none of these' "
                          f"if it has none")
            continue
        items[vid]["content"] = values
        items[vid]["watchedBy"] = {"by": who, "on": today}
        changed.append(vid)
    return changed, report
